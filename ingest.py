"""环境法规问答助手 · 建库脚本

用法：python ingest.py
将 data/raw/ 下的法规文件（.docx / .pdf）解析、按条文切块、向量化后
写入 Chroma 向量库（data/vector_db/）。重复运行会清空重建，可安全重复执行。

切分策略（条文语义切分）：
  - "第X章" 行记录为当前章标题，并拼进该章每条条文作上下文
  - "第X条" 行开启新条文，其后非章非条段落归入当前条文
  - 超长条文（>500 字）内部按字符二次切分
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import chromadb
from langchain_text_splitters import RecursiveCharacterTextSplitter
from sentence_transformers import SentenceTransformer

# ---------- 配置 ----------
BASE_DIR = Path(__file__).resolve().parent
RAW_DIR = BASE_DIR / "data" / "raw"
VECTOR_DB_DIR = BASE_DIR / "data" / "vector_db"
COLLECTION_NAME = "env_law"
MODEL_NAME = "BAAI/bge-small-zh-v1.5"
CHUNK_SIZE = 500          # 超长条文二次切分的单块上限
CHUNK_OVERLAP = 100       # 二次切分的重叠字符数
BATCH_SIZE = 100          # 写入 Chroma 的批量大小

# 条文/章节行首识别："第三条""第一百零五条"等
ARTICLE_RE = re.compile(r"^第[一二三四五六七八九十百零0-9]+条")
CHAPTER_RE = re.compile(r"^第[一二三四五六七八九十百0-9]+章")


def extract_text(path: Path) -> list[tuple[str, int]]:
    """按扩展名分流解析，返回 [(文本, 页码)]。docx 无页码概念，记 0。"""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".pdf":
        return _extract_pdf(path)
    print(f"[跳过] 不支持的格式：{path.name}")
    return []


def _extract_docx(path: Path) -> list[tuple[str, int]]:
    import docx
    doc = docx.Document(str(path))
    texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    # 表格内容也并入，避免丢正文
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    texts.append(cell_text)
    return [(t, 0) for t in texts]


def _extract_pdf(path: Path) -> list[tuple[str, int]]:
    import fitz
    texts: list[tuple[str, int]] = []
    with fitz.open(str(path)) as doc:
        for page_num, page in enumerate(doc, start=1):
            text = page.get_text().strip()
            if text:
                texts.append((text, page_num))
            else:
                print(f"[警告] {path.name} 第 {page_num} 页无文字，已跳过")
    return texts


def parse_articles(paragraphs: list[tuple[str, int]]) -> list[dict]:
    """把段落流组织成条文列表。

    规则：
    - "第X章" 行：记录为当前章标题
    - "第X条" 行：开启新条文
    - 其余行：归入当前条文（处理长条文跨段、款/项列举）
    - 条文之前的杂项（标题、日期、目录）丢弃

    返回：[{"chapter": 章标题|None, "num": 条文号|None, "text": 全文, "page": int}]
    """
    articles: list[dict] = []
    current_chapter: str | None = None
    current: dict | None = None
    for text, page in paragraphs:
        if CHAPTER_RE.match(text):
            current_chapter = text
            continue
        if ARTICLE_RE.match(text):
            if current is not None:
                articles.append(current)
            num = ARTICLE_RE.match(text).group()
            current = {"chapter": current_chapter, "num": num, "text": text, "page": page}
            continue
        if current is not None:
            current["text"] += "\n" + text
    if current is not None:
        articles.append(current)
    return articles


def split_articles(articles: list[dict]) -> list[dict]:
    """把条文切成入库块：章标题拼进每条，超长条文内部二次切分。"""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP
    )
    chunks: list[dict] = []
    for art in articles:
        body = art["text"]
        # 章标题作为上下文拼到条文前
        doc_text = f'{art["chapter"]}\n{body}' if art["chapter"] else body
        if len(doc_text) <= CHUNK_SIZE:
            chunks.append({**art, "text": doc_text})
            continue
        # 超长条文：内部按字符二次切分，保留同样的出处信息
        for part in splitter.split_text(doc_text):
            chunks.append({**art, "text": part})
    return chunks


def main() -> None:
    files = sorted(
        p for p in RAW_DIR.rglob("*")
        if p.is_file() and p.suffix.lower() in (".docx", ".pdf")
    )
    if not files:
        print(f"[阻塞] {RAW_DIR} 下没有 .docx/.pdf 文件，请先放入法规文件。")
        sys.exit(1)

    print(f"[加载] 向量模型 {MODEL_NAME}（首次运行会联网下载）...")
    try:
        model = SentenceTransformer(MODEL_NAME)
    except Exception as exc:
        print(f"[阻塞] 模型加载失败：{exc}")
        print("若为网络原因，请先执行：")
        print("  set HF_ENDPOINT=https://hf-mirror.com")
        print("然后再重新运行本脚本。")
        sys.exit(1)

    documents: list[str] = []
    metadatas: list[dict] = []
    ids: list[str] = []
    file_count = 0
    chunk_count = 0
    global_idx = 0
    for path in files:
        paragraphs = extract_text(path)
        if not paragraphs:
            print(f"[警告] {path.name} 未提取到任何文本，已跳过")
            continue
        articles = parse_articles(paragraphs)
        chunks = split_articles(articles)
        file_count += 1
        chunk_count += len(chunks)
        for chunk in chunks:
            documents.append(chunk["text"])
            metadatas.append({
                "source": path.name,
                "page": chunk["page"],
                "num": chunk["num"] or "",
                "chapter": chunk["chapter"] or "",
            })
            ids.append(f"doc-{global_idx}")
            global_idx += 1
        print(f"[完成] {path.name}：{len(articles)} 条条文 → {len(chunks)} 块")

    print(f"[向量化] 共 {len(documents)} 块，生成 {model.get_sentence_embedding_dimension()} 维向量...")
    embeddings = model.encode(documents, normalize_embeddings=True)

    VECTOR_DB_DIR.mkdir(parents=True, exist_ok=True)
    client = chromadb.PersistentClient(path=str(VECTOR_DB_DIR))
    try:
        client.delete_collection(COLLECTION_NAME)
        print("[重建] 已清空旧 collection")
    except Exception:
        pass
    collection = client.get_or_create_collection(COLLECTION_NAME)

    print("[写入] 写入 Chroma...")
    for start in range(0, len(documents), BATCH_SIZE):
        end = start + BATCH_SIZE
        collection.add(
            ids=ids[start:end],
            documents=documents[start:end],
            metadatas=metadatas[start:end],
            embeddings=embeddings[start:end].tolist(),
        )

    print("\n===== 建库完成 =====")
    print(f"处理文件数：{file_count}")
    print(f"总块数：{chunk_count}")
    print(f"平均每文件块数：{chunk_count / file_count if file_count else 0:.1f}")
    print(f"向量库位置：{VECTOR_DB_DIR}")


if __name__ == "__main__":
    main()
